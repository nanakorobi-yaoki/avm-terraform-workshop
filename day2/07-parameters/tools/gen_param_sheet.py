#!/usr/bin/env python3
"""
納品用パラメータシート (Excel) 自動生成ツール
=============================================================================
Excel を「入力」ではなく「成果物」として扱うための仕組みです。

  従来: Excel (手入力) ──→ 人手転記 ──→ Terraform     ※ 転記ミス・二重管理
  提案: YAML (SSOT) ──→ Terraform ──→ Excel (自動生成) ※ 常に実態と一致

生成内容:
  Sheet1 「メタ情報」     : 生成日時・Git コミット・依存バージョン・入力の SHA-256
  Sheet2 「パラメータ一覧」 : YAML の値 + JSON Schema の説明
  Sheet3 「リソース一覧」   : terraform plan の結果（実際に作られるもの）

使い方:
    pip install openpyxl pyyaml
    terraform show -json tfplan > tfplan.json     # 任意
    python tools/gen_param_sheet.py \
        --params params/contoso-prod.yaml \
        --schema schema/params.schema.json \
        --plan   tfplan.json \
        --output dist/contoso-prod-パラメータシート.xlsx

手編集検知（CI）:
    # xlsx は ZIP バイナリで内部に生成日時を含むため、バイト比較できません。
    # 決定的な正規化 JSON を介して差分を検証します。
    python tools/gen_param_sheet.py ... --format json --output dist/params.normalized.json
    git diff --exit-code dist/params.normalized.json
=============================================================================
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    import yaml
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("依存パッケージが不足しています: pip install openpyxl pyyaml")


HEADER_FILL = PatternFill("solid", fgColor="0078D4")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def flatten(data: Any, schema: dict, prefix: str = "") -> list[dict]:
    """YAML を Schema の説明とマージしつつ 1 行 1 項目に平坦化する。"""
    rows: list[dict] = []
    props = schema.get("properties", {}) if isinstance(schema, dict) else {}

    if isinstance(data, dict):
        for key, value in data.items():
            path = f"{prefix}.{key}" if prefix else key
            sub_schema = props.get(key)
            if sub_schema is None:
                addl = schema.get("additionalProperties")
                sub_schema = addl if isinstance(addl, dict) else {}

            if isinstance(value, (dict, list)) and value:
                if isinstance(value, list) and all(
                    not isinstance(i, (dict, list)) for i in value
                ):
                    rows.append(_row(path, value, sub_schema))
                else:
                    child_schema = (
                        sub_schema.get("items", {})
                        if isinstance(value, list)
                        else sub_schema
                    )
                    rows.extend(flatten(value, child_schema, path))
            else:
                rows.append(_row(path, value, sub_schema))

    elif isinstance(data, list):
        for idx, item in enumerate(data):
            rows.extend(flatten(item, schema, f"{prefix}[{idx}]"))

    return rows


def _row(path: str, value: Any, schema: dict) -> dict:
    if isinstance(value, list):
        display = ", ".join(str(v) for v in value)
    elif isinstance(value, bool):
        display = "有効" if value else "無効"
    else:
        display = str(value)

    constraints = []
    if "enum" in schema:
        constraints.append("許可値: " + " / ".join(map(str, schema["enum"])))
    if "pattern" in schema:
        constraints.append(f"形式: {schema['pattern']}")
    if "minimum" in schema or "maximum" in schema:
        constraints.append(
            f"範囲: {schema.get('minimum', '-')} 〜 {schema.get('maximum', '-')}"
        )
    if schema.get("format"):
        constraints.append(f"形式: {schema['format']}")

    return {
        "path": path,
        "title": schema.get("title", ""),
        "value": display,
        "type": schema.get("type", ""),
        "description": schema.get("description", ""),
        "constraints": " / ".join(constraints),
    }


def style_sheet(ws, headers: list[str], widths: list[int], title: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"生成日時: {datetime.now():%Y-%m-%d %H:%M:%S}（自動生成 — 直接編集しないでください）"
    ws["A2"].font = Font(size=9, color="808080")

    for col, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A5"


def build_param_sheet(wb: Workbook, params: dict, schema: dict) -> None:
    ws = wb.active
    ws.title = "パラメータ一覧"
    headers = ["項目", "項目名", "設定値", "型", "説明", "制約"]
    style_sheet(ws, headers, [30, 24, 30, 10, 46, 34], "環境構築パラメータシート")

    for i, row in enumerate(flatten(params, schema), start=5):
        for col, key in enumerate(
            ["path", "title", "value", "type", "description", "constraints"], start=1
        ):
            cell = ws.cell(row=i, column=col, value=row[key])
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col in (5, 6))


def build_resource_sheet(wb: Workbook, plan_path: Path | None) -> None:
    ws = wb.create_sheet("リソース一覧")
    headers = ["アドレス", "リソース種別", "リソース名", "操作"]
    style_sheet(ws, headers, [46, 40, 34, 12], "作成されるリソース一覧 (terraform plan より)")

    if not plan_path or not plan_path.exists():
        ws.cell(row=5, column=1, value="(tfplan.json が指定されていません)")
        return

    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    action_ja = {
        "create": "新規作成",
        "update": "更新",
        "delete": "削除",
        "no-op": "変更なし",
        "read": "参照",
    }

    row_no = 5
    for change in plan.get("resource_changes", []):
        actions = change.get("change", {}).get("actions", [])
        if actions == ["no-op"]:
            continue
        after = change.get("change", {}).get("after") or {}
        values = [
            change.get("address", ""),
            change.get("type", ""),
            after.get("name", change.get("name", "")),
            " / ".join(action_ja.get(a, a) for a in actions),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_no, column=col, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top")
        row_no += 1


def _git(*args: str) -> str:
    """git コマンドを実行する。失敗時は空文字列を返す。"""
    try:
        out = subprocess.run(
            ["git", *args],
            capture_output=True,
            text=True,
            timeout=5,
            check=False,
        )
        return out.stdout.strip() if out.returncode == 0 else ""
    except (OSError, subprocess.SubprocessError):
        return ""


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _module_versions(module_dir: Path) -> list[str]:
    """.terraform/modules/modules.json から依存モジュールとバージョンを拾う。"""
    manifest = module_dir / ".terraform" / "modules" / "modules.json"
    if not manifest.exists():
        return []
    try:
        data = json.loads(manifest.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []

    found = {
        f"{m['Source']} {m.get('Version', '')}".strip()
        for m in data.get("Modules", [])
        if m.get("Source", "").startswith("registry.terraform.io/")
    }
    return sorted(found)


def collect_provenance(params_path: Path, schema_path: Path) -> list[tuple[str, str]]:
    """納品物に刻印するトレーサビリティ情報を集める。

    CI 上では GitHub Actions の環境変数を優先し、手元実行時は git に落とす。
    """
    sha = os.environ.get("GITHUB_SHA") or _git("rev-parse", "HEAD")
    dirty = _git("status", "--porcelain")
    server = os.environ.get("GITHUB_SERVER_URL", "https://github.com")
    repo = os.environ.get("GITHUB_REPOSITORY", "")
    run_id = os.environ.get("GITHUB_RUN_ID", "")
    run_url = f"{server}/{repo}/actions/runs/{run_id}" if repo and run_id else ""

    rows: list[tuple[str, str]] = [
        ("生成日時 (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")),
        ("Git コミット", sha or "(git リポジトリ外)"),
        ("リリースタグ", _git("describe", "--tags", "--always", "--dirty") or "-"),
        ("ブランチ", os.environ.get("GITHUB_REF_NAME") or _git("rev-parse", "--abbrev-ref", "HEAD") or "-"),
        (
            "作業ツリーの状態",
            "未コミットの変更あり ⚠️" if dirty else "クリーン（コミット済み）",
        ),
        ("生成者", os.environ.get("GITHUB_ACTOR") or "(手動実行)"),
        ("CI 実行 URL", run_url or "(手動実行のため CI 記録なし) ⚠️"),
        ("パラメータファイル", str(params_path)),
        ("パラメータの SHA-256", _sha256(params_path)),
        ("スキーマの SHA-256", _sha256(schema_path)),
    ]

    modules = _module_versions(params_path.parent.parent)
    rows.append(
        ("依存モジュール", "\n".join(modules) if modules else "(terraform init 未実行)")
    )
    return rows


def build_meta_sheet(wb: Workbook, provenance: list[tuple[str, str]]) -> None:
    """『この Excel はいつ・どのコードから生成されたか』を示すシート。"""
    ws = wb.create_sheet("メタ情報", 0)
    style_sheet(ws, ["項目", "値"], [30, 96], "納品物メタ情報 (トレーサビリティ)")

    for i, (label, value) in enumerate(provenance, start=5):
        key_cell = ws.cell(row=i, column=1, value=label)
        key_cell.font = Font(bold=True)
        key_cell.border = BORDER
        key_cell.alignment = Alignment(vertical="top")

        val_cell = ws.cell(row=i, column=2, value=value)
        val_cell.border = BORDER
        val_cell.alignment = Alignment(vertical="top", wrap_text=True)


def normalized_payload(params: dict, schema: dict, plan_path: Path | None) -> dict:
    """CI の手編集検知に使う、決定的な正規化表現を作る。

    ★ 生成日時や Git ハッシュなど『実行ごとに変わる値』は意図的に含めない。
      含めると毎回差分が出てしまい、改ざん検知として機能しなくなる。
    """
    payload: dict[str, Any] = {
        "parameters": flatten(params, schema),
        "resources": [],
    }

    if plan_path and plan_path.exists():
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        payload["resources"] = sorted(
            (
                {
                    "address": c.get("address", ""),
                    "type": c.get("type", ""),
                    "actions": c.get("change", {}).get("actions", []),
                }
                for c in plan.get("resource_changes", [])
                if c.get("change", {}).get("actions") != ["no-op"]
            ),
            key=lambda r: r["address"],
        )

    return payload


def main() -> int:
    ap = argparse.ArgumentParser(description="納品用パラメータシートを生成します")
    ap.add_argument("--params", required=True, type=Path, help="パラメータ YAML")
    ap.add_argument("--schema", required=True, type=Path, help="JSON Schema")
    ap.add_argument("--plan", type=Path, help="terraform show -json の出力")
    ap.add_argument("--output", required=True, type=Path, help="出力パス")
    ap.add_argument(
        "--format",
        choices=["xlsx", "json"],
        default="xlsx",
        help="xlsx=納品用 / json=CI の手編集検知用（決定的な正規化表現）",
    )
    args = ap.parse_args()

    params = yaml.safe_load(args.params.read_text(encoding="utf-8"))
    schema = json.loads(args.schema.read_text(encoding="utf-8"))
    args.output.parent.mkdir(parents=True, exist_ok=True)

    if args.format == "json":
        payload = normalized_payload(params, schema, args.plan)
        args.output.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        print(f"正規化 JSON を生成しました: {args.output}")
        return 0

    wb = Workbook()
    build_param_sheet(wb, params, schema)
    build_resource_sheet(wb, args.plan)
    build_meta_sheet(wb, collect_provenance(args.params, args.schema))

    wb.save(args.output)
    print(f"生成しました: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
